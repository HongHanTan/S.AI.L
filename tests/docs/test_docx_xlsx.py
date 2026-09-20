import io
import zipfile

import pytest
from docx import Document

from sdoc.docs.docx_ import read_docx
from sdoc.docs.xlsx import read_xlsx

SHEET_XML = """<?xml version="1.0"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData>
<row r="1"><c r="A1" t="inlineStr"><is><t>APRIL FAR EAST (M) SDN BHD</t></is></c></row>
<row r="3"><c r="A3" t="inlineStr"><is><t>BL INSTRUCTION</t></is></c>
           <c r="B3" t="inlineStr"><is><t>3815798123</t></is></c></row>
<row r="4"><c r="A4" t="inlineStr"><is><t>CONSIGNEE</t></is></c>
           <c r="B4" t="inlineStr"><is><t>BALL &amp; DOGGETT | 43-45 METRO RD</t></is></c></row>
<row r="5"><c r="A5" t="inlineStr"><is><t>Load Port</t></is></c>
           <c r="B5" t="inlineStr"><is><t>SINGAPORE</t></is></c></row>
</sheetData></worksheet>"""


def _xlsx_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", SHEET_XML)
        z.writestr("[Content_Types].xml", "<Types/>")
    return buf.getvalue()


def test_xlsx_emits_label_colon_value_lines():
    lines = read_xlsx(_xlsx_bytes())
    assert "CONSIGNEE: BALL & DOGGETT | 43-45 METRO RD" in lines
    assert "Load Port: SINGAPORE" in lines


def test_xlsx_header_row_appears_in_first_lines():
    lines = read_xlsx(_xlsx_bytes())
    non_empty = [x for x in lines if x.strip()]
    assert any("BL INSTRUCTION" in x for x in non_empty[:3])


def test_xlsx_decodes_html_entities():
    assert "&amp;" not in "\n".join(read_xlsx(_xlsx_bytes()))


def test_docx_reads_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("BILL OF LADING (DRAFT)")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Load Port"
    table.cell(0, 1).text = "SINGAPORE"
    buf = io.BytesIO()
    doc.save(buf)

    lines = read_docx(buf.getvalue())
    assert "BILL OF LADING (DRAFT)" in lines
    assert "Load Port: SINGAPORE" in lines
