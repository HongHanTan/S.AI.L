"""Read .docx paragraphs and tables into 'Label: value' lines."""
import io

from docx import Document


def read_docx(raw: bytes) -> list[str]:
    doc = Document(io.BytesIO(raw))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) >= 2 and cells[0] and cells[1]:
                lines.append(f"{cells[0]}: {cells[1]}")
            elif cells and cells[0]:
                lines.append(cells[0])
    return lines
